#!/usr/bin/env python3
"""
Render a deconstruction report (.md) into a .docx laid out like the user's
existing COMPETITOR_ADS documents: headings, bold beat lines, real Word tables.

    python3 build_docx.py report.md --out COMPETITOR_ADS_2.docx

Requires python-docx (pip install python-docx).
"""
import argparse, os, re, sys

try:
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    sys.exit("python-docx not installed. Run: pip install python-docx")

INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`)")

def add_runs(par, text):
    """Render **bold**, *italic* and `code` inside a paragraph."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(9)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            par.add_run(part[1:-1]).italic = True
        else:
            par.add_run(part)

def is_row(l):    return l.strip().startswith("|") and l.strip().count("|") >= 2
def is_sep(l):    return bool(re.fullmatch(r"\|[\s:\-|]+\|", l.strip()))
def cells(l):     return [c.strip() for c in l.strip().strip("|").split("|")]

def add_table(doc, rows):
    header, body = rows[0], rows[1:]
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.autofit = True
    for i, h in enumerate(header):
        p = t.rows[0].cells[i].paragraphs[0]
        add_runs(p, h)
        for r in p.runs:
            r.bold = True; r.font.size = Pt(9)
    for row in body:
        rc = t.add_row().cells
        for i, c in enumerate(row[:len(header)]):
            p = rc[i].paragraphs[0]
            add_runs(p, c)
            for r in p.runs:
                r.font.size = Pt(9)
    doc.add_paragraph()

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("markdown")
    ap.add_argument("--out", default=None)
    ap.add_argument("--title", default=None)
    a = ap.parse_args()

    src = open(a.markdown, encoding="utf-8").read()
    out = a.out or os.path.splitext(a.markdown)[0] + ".docx"

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"; st.font.size = Pt(10.5)
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(0.8)

    if a.title:
        h = doc.add_heading(a.title, level=0)

    lines, i, buf = src.splitlines(), 0, []
    while i < len(lines):
        line = lines[i]

        if is_row(line):
            tbl = []
            while i < len(lines) and is_row(lines[i]):
                if not is_sep(lines[i]):
                    tbl.append(cells(lines[i]))
                i += 1
            if tbl:
                add_table(doc, tbl)
            continue

        if line.strip().startswith("```"):
            i += 1
            code = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i]); i += 1
            i += 1
            p = doc.add_paragraph()
            r = p.add_run("\n".join(code))
            r.font.name = "Consolas"; r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
            continue

        m = re.match(r"^(#{1,4})\s+(.*)$", line)
        if m:
            doc.add_heading(re.sub(r"\*\*", "", m.group(2)),
                            level=min(len(m.group(1)), 4))
            i += 1; continue

        if re.fullmatch(r"\s*-{3,}\s*", line):
            doc.add_paragraph()
            i += 1; continue

        bm = re.match(r"^\s*[-*]\s+(.*)$", line)
        if bm:
            add_runs(doc.add_paragraph(style="List Bullet"), bm.group(1))
            i += 1; continue

        bq = re.match(r"^>\s?(.*)$", line)
        if bq:
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.35)
            add_runs(p, bq.group(1))
            for r in p.runs: r.italic = True
            i += 1; continue

        if line.strip():
            add_runs(doc.add_paragraph(), line.strip())
        i += 1

    doc.save(out)
    print(f"Wrote {out}")

if __name__ == "__main__":
    main()
